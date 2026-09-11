#!/usr/bin/env python3
"""
Planilha Embriologia Loader - Load all Excel files from year subfolders to bronze layer
Reads all Excel files from planilha_embriologia/data_input/YYYY/ folders and loads them to bronze tables.
Supports years 2021 through 2026, including FRESH, FET, FOT, RECEP, and consolidated historical sheets.
Fully overwrites bronze tables on each run.
"""

import logging
import pandas as pd
import duckdb
from datetime import datetime
import os
import glob
import re
import openpyxl
import unicodedata

# Setup logging
LOGS_DIR = os.path.join(os.path.dirname(__file__), 'logs')
os.makedirs(LOGS_DIR, exist_ok=True)
timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
script_name = os.path.splitext(os.path.basename(__file__))[0]
LOG_PATH = os.path.join(LOGS_DIR, f'{script_name}_{timestamp}.log')

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s %(levelname)s %(message)s',
    handlers=[
        logging.FileHandler(LOG_PATH),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# Configuration
DUCKDB_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), 'database', 'huntington_data_lake.duckdb')
DATA_INPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'data_input')
CLINICAL_SHEETS_ADDITIONAL = [
    'DOADORAS',
    'FP (cong ovulos e tecidos)',
    'FP (cong ovulos e tecido)',
    'FP',
    'FP (cong de Semen)',
    'IIU'
]
SHEETS_TO_LOAD = ['FRESH', 'FET', 'FOT', 'RECEP'] + CLINICAL_SHEETS_ADDITIONAL

# Year-specific configurations
YEAR_CONFIGS = {
    'DEFAULT': {
        'sheets': ['FRESH', 'FET', 'FOT', 'RECEP'] + CLINICAL_SHEETS_ADDITIONAL,
        'header': 1
    },
    '2021': {
        'sheets': ['ANUAL JAN-DEZ CERTO', 'TOTAL 2021', 'TOTAL'],
        'header': 1
    },
    '2022': {
        'sheets': ['TOTAL', 'TOTAL 2022', '2022', 'FRESH', 'FET', 'FOT', 'RECEP', 'FIV', 'TEC', 'IIU'],
        'header': 1
    },
    '2023': {
        'sheets': ['FRESH', 'FET', 'FOT', 'RECEP', 'Total 2023 Nova ', 'Total 2023 Nova', 'GERAL 2023', 'FIV', 'TEC'] + CLINICAL_SHEETS_ADDITIONAL,
        'header': 1
    },
    '2024': {
        'sheets': ['FRESH', 'FET', 'FOT', 'RECEP'] + CLINICAL_SHEETS_ADDITIONAL,
        'header': 1
    },
    '2025': {
        'sheets': ['FRESH', 'FET', 'FOT', 'RECEP'] + CLINICAL_SHEETS_ADDITIONAL,
        'header': 1
    },
    '2026': {
        'sheets': ['FRESH', 'FET', 'FOT', 'RECEP'] + CLINICAL_SHEETS_ADDITIONAL,
        'header': 1
    }
}

def get_duckdb_connection():
    """Create DuckDB connection"""
    try:
        logger.info(f"Attempting to connect to DuckDB at: {DUCKDB_PATH}")
        con = duckdb.connect(DUCKDB_PATH)
        logger.info("DuckDB connection successful")
        return con
    except Exception as e:
        logger.error(f"Failed to connect to DuckDB: {e}")
        raise

def get_all_excel_files():
    """Get all Excel files from year subfolders (YYYY), ignoring further subfolders and duplicates"""
    excel_files = []
    
    if not os.path.exists(DATA_INPUT_DIR):
        raise FileNotFoundError(f"Data input directory not found: {DATA_INPUT_DIR}")
    
    # List all items in data_input directory
    for item in sorted(os.listdir(DATA_INPUT_DIR)):
        item_path = os.path.join(DATA_INPUT_DIR, item)
        
        # Check if it's a directory and looks like a year (4 digits)
        if os.path.isdir(item_path) and item.isdigit() and len(item) == 4:
            logger.info(f"Scanning year folder: {item}")
            
            # Get all Excel files directly in this year folder
            year_excel_files = glob.glob(os.path.join(item_path, "*.xlsx"))
            
            # Filter out temporary Excel files (files starting with ~$)
            year_excel_files = [f for f in year_excel_files if not os.path.basename(f).startswith('~$')]
            
            # Deduplicate: If both 'CASOS YYYY IBIRA.xlsx' and 'CASOS YYYY IBI.xlsx' exist, drop 'IBI.xlsx'
            base_names = [os.path.basename(f).upper() for f in year_excel_files]
            has_ibira = any('IBIRA' in b for b in base_names)
            
            filtered_files = []
            for f in year_excel_files:
                bname = os.path.basename(f).upper()
                if has_ibira and (' IBI.XLSX' in bname or ' IBI ' in bname):
                    logger.info(f"Skipping duplicate file: {os.path.basename(f)} (using IBIRA instead)")
                    continue
                # Exclude 2022 BSB
                if item == '2022' and 'BSB' in bname:
                    logger.info(f"Skipping excluded file: {os.path.basename(f)} (2022 BSB excluded from ingestion)")
                    continue
                filtered_files.append(f)
            
            excel_files.extend(filtered_files)
            logger.info(f"Found {len(filtered_files)} Excel file(s) in {item}/")
    
    if not excel_files:
        raise FileNotFoundError(f"No Excel files found in year subfolders of {DATA_INPUT_DIR}")
    
    logger.info(f"Total Excel files found: {len(excel_files)}")
    return sorted(excel_files)

def detect_header_row(file_path, sheet_name, max_rows=15):
    """
    Attempt to detect the header row by looking for 'PIN', 'PRONTUARIO', 'PACIENTE', 'DATA', etc.
    in the first max_rows rows using openpyxl. Returns the 0-indexed row number.
    """
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        if sheet_name not in wb.sheetnames:
            wb.close()
            return None
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True, max_row=max_rows))
        wb.close()
        for idx, row in enumerate(rows):
            cells = [unicodedata.normalize('NFD', str(c)).upper() for c in row if c is not None]
            has_paciente = any('PACIENTE' in c or 'NOME' in c for c in cells)
            has_pin = any('PRONTUARIO' in c or 'PIN' in c or 'CONTROLE' in c or 'CHART' in c for c in cells)
            has_data = any('DATA' in c or 'DIA' in c or 'DATE' in c for c in cells)
            has_tipo = any('TIPO' in c or 'PROCED' in c or 'TRATAMENTO' in c or 'PROCEDURE' in c for c in cells)
            if (has_paciente and has_data) or (has_pin and has_data) or (has_tipo and has_data):
                return idx
    except Exception as e:
        logger.warning(f"Header detection failed for {file_path} [{sheet_name}]: {e}")
    return None

def generate_table_name(file_path, sheet_name):
    """Generate standardized table name: planilha_{year}_{location}_{sheet}"""
    try:
        year = os.path.basename(os.path.dirname(file_path))
        filename = os.path.basename(file_path)
        name_no_ext = os.path.splitext(filename)[0]
        
        # Clean up the name
        clean_name = name_no_ext.lower().replace('casos', '').replace(year, '')
        clean_name = re.sub(r'[^a-z0-9]', '_', clean_name)
        clean_name = re.sub(r'_+', '_', clean_name).strip('_')
        
        # Construct table name
        sheet_clean = re.sub(r'[^a-z0-9]', '_', sheet_name.lower())
        sheet_clean = re.sub(r'_+', '_', sheet_clean).strip('_')
        
        table_name = f"planilha_{year}_{clean_name}_{sheet_clean}"
        safe_table_name = re.sub(r'[^a-zA-Z0-9_]', '_', table_name).lower()
        return safe_table_name
    except Exception as e:
        logger.warning(f"Could not generate standard name, falling back to safe filename: {e}")
        base = os.path.splitext(os.path.basename(file_path))[0]
        return f"planilha_{base}_{sheet_name}".lower()

def create_bronze_table(con, table_name, columns):
    """Create a bronze table with all columns as VARCHAR/TEXT plus metadata"""
    try:
        con.execute("CREATE SCHEMA IF NOT EXISTS bronze")
        con.execute(f"DROP TABLE IF EXISTS bronze.{table_name}")
        
        col_defs = []
        used_names = set()
        
        for col in columns:
            col_str = str(col) if pd.notna(col) and str(col).strip() != '' else 'unnamed_col'
            counter = 0
            unique_name = col_str
            while unique_name.lower() in [n.lower() for n in used_names]:
                counter += 1
                unique_name = f"{col_str}_{counter}"
            used_names.add(unique_name)
            col_defs.append(f'"{unique_name}" VARCHAR')
        
        # Add metadata columns
        col_defs.extend([
            'line_number INTEGER',
            'extraction_timestamp VARCHAR',
            'file_name VARCHAR',
            'sheet_name VARCHAR'
        ])
        
        create_sql = f"""
        CREATE TABLE bronze.{table_name} (
            {', '.join(col_defs)}
        )
        """
        con.execute(create_sql)
        logger.info(f"Table bronze.{table_name} created successfully with {len(columns)} data columns")
    except Exception as e:
        logger.error(f"Error creating bronze table {table_name}: {e}")
        raise

def process_ssa_2022_file(file_path, con):
    """Special handler for CASOS 2022 SSA.xlsx: unions monthly FIV and TEC sheets"""
    file_name = os.path.basename(file_path)
    year = "2022"
    total_loaded = 0
    
    try:
        xl_file = pd.ExcelFile(file_path, engine='openpyxl')
        sheets = xl_file.sheet_names
        
        # 1. FIV sheets (Fresh)
        fiv_sheets = [s for s in sheets if s.upper().startswith('FIV ')]
        if fiv_sheets:
            logger.info(f"Consolidating {len(fiv_sheets)} FIV monthly sheets for {file_name}...")
            fiv_dfs = []
            for s in fiv_sheets:
                header_row = detect_header_row(file_path, s, max_rows=15)
                if header_row is None:
                    header_row = 0 if 'AGOSTO' in s.upper() else 6
                df_s = pd.read_excel(file_path, sheet_name=s, header=header_row, dtype=str, engine='openpyxl')
                df_s['sheet_name'] = s
                fiv_dfs.append(df_s)
            
            df_fiv = pd.concat(fiv_dfs, ignore_index=True)
            table_name = "planilha_2022_ssa_fiv"
            total_loaded += insert_dataframe_to_bronze(con, df_fiv, table_name, file_name, "FIV_MONTHLY_CONSOLIDATED")
        
        # 2. TEC sheets (FET)
        tec_sheets = [s for s in sheets if s.upper().startswith('TEC ')]
        if tec_sheets:
            logger.info(f"Consolidating {len(tec_sheets)} TEC monthly sheets for {file_name}...")
            tec_dfs = []
            for s in tec_sheets:
                header_row = detect_header_row(file_path, s, max_rows=15)
                if header_row is None:
                    header_row = 4
                df_s = pd.read_excel(file_path, sheet_name=s, header=header_row, dtype=str, engine='openpyxl')
                df_s['sheet_name'] = s
                tec_dfs.append(df_s)
            
            df_tec = pd.concat(tec_dfs, ignore_index=True)
            table_name = "planilha_2022_ssa_tec"
            total_loaded += insert_dataframe_to_bronze(con, df_tec, table_name, file_name, "TEC_MONTHLY_CONSOLIDATED")
            
        # 3. IIU sheet
        if 'IIU' in sheets:
            header_row = 5
            logger.info(f"Loading IIU sheet for {file_name} with explicit header row {header_row}...")
            df_iiu = pd.read_excel(file_path, sheet_name='IIU', header=header_row, dtype=str, engine='openpyxl')
            table_name = "planilha_2022_ssa_iiu"
            total_loaded += insert_dataframe_to_bronze(con, df_iiu, table_name, file_name, "IIU")
            
    except Exception as e:
        logger.error(f"Error processing SSA 2022 monthly consolidation: {e}")
    
    return total_loaded

def insert_dataframe_to_bronze(con, df, table_name, file_name, sheet_name):
    """Insert a prepared DataFrame into a bronze table"""
    if len(df) == 0:
        logger.warning(f"No data to insert for {table_name}")
        return 0
    
    original_columns = [c for c in df.columns if c not in ['line_number', 'extraction_timestamp', 'file_name', 'sheet_name']]
    create_bronze_table(con, table_name, original_columns)
    
    used_names = set()
    column_mapping = {}
    unique_columns = []
    
    for col in original_columns:
        original_col = str(col) if pd.notna(col) and str(col).strip() != '' else 'unnamed_col'
        counter = 0
        unique_col = original_col
        while unique_col.lower() in [n.lower() for n in used_names]:
            counter += 1
            unique_col = f"{original_col}_{counter}"
        column_mapping[col] = unique_col
        unique_columns.append(unique_col)
        used_names.add(unique_col)
    
    df_renamed = df[original_columns].rename(columns=column_mapping)
    for col in df_renamed.columns:
        df_renamed[col] = df_renamed[col].replace('', None)
    
    extraction_timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    df_renamed['line_number'] = df_renamed.index
    df_renamed['extraction_timestamp'] = extraction_timestamp
    df_renamed['file_name'] = file_name
    if 'sheet_name' in df.columns:
        df_renamed['sheet_name'] = df['sheet_name']
    else:
        df_renamed['sheet_name'] = sheet_name
    
    con.execute(f"INSERT INTO bronze.{table_name} SELECT * FROM df_renamed")
    logger.info(f"Successfully inserted {len(df_renamed)} rows into bronze.{table_name}")
    return len(df_renamed)

def process_excel_file(file_path, con):
    """Process a single Excel file and load data from all configured sheets"""
    file_name = os.path.basename(file_path)
    total_loaded = 0
    year = os.path.basename(os.path.dirname(file_path))
    
    # Check for SSA 2022 monthly layout
    if year == '2022' and 'SSA' in file_name.upper():
        return process_ssa_2022_file(file_path, con)
    
    config = YEAR_CONFIGS.get(year, YEAR_CONFIGS['DEFAULT'])
    sheets_to_try = config['sheets']
    default_header_row = config['header']
    
    try:
        xl_file = pd.ExcelFile(file_path, engine='openpyxl')
        actual_sheets = xl_file.sheet_names
        actual_sheets_lower = [s.lower() for s in actual_sheets]
    except Exception as e:
        logger.error(f"Could not read excel file {file_name}: {e}")
        return 0

    # Determine which sheets to process
    sheets_to_process = []
    for sheet_pattern in sheets_to_try:
        try:
            idx = actual_sheets_lower.index(sheet_pattern.lower())
            actual_name = actual_sheets[idx]
            if actual_name not in sheets_to_process:
                sheets_to_process.append(actual_name)
        except ValueError:
            continue

    # For CASOS 2023 IBIRA.xlsx, prioritize Total 2023 Nova over Total 2023
    if year == '2023' and 'IBIRA' in file_name.upper():
        if any('nova' in s.lower() for s in sheets_to_process):
            sheets_to_process = [s for s in sheets_to_process if 'nova' in s.lower()]

    if not sheets_to_process:
        logger.warning(f"None of configured sheets {sheets_to_try} found in {file_name} (Available: {actual_sheets})")
        return 0

    for sheet in sheets_to_process:
        table_name = generate_table_name(file_path, sheet)
        
        # Header row detection
        if 'SSA' in file_name.upper() and sheet.upper() == 'IIU':
            if year in ['2022', '2023']:
                actual_header_row = 5
            elif year in ['2025', '2026']:
                actual_header_row = 0
            else:
                actual_header_row = 1
            logger.info(f"Using explicit header row {actual_header_row} for SSA IIU {file_name} [{sheet}]")
        else:
            actual_header_row = default_header_row
            detected = detect_header_row(file_path, sheet)
            if detected is not None:
                actual_header_row = detected
                logger.info(f"Detected header for {file_name} [{sheet}] at row {actual_header_row}")
            else:
                logger.info(f"Using default header row {actual_header_row} for {file_name} [{sheet}]")
        
        logger.info(f"Processing: {file_name} [{sheet}] (Year: {year}, Header Row: {actual_header_row}) -> bronze.{table_name}")
        
        try:
            df = pd.read_excel(
                file_path, 
                sheet_name=sheet,
                engine='openpyxl',
                header=actual_header_row,
                dtype=str
            )
            
            rows_loaded = insert_dataframe_to_bronze(con, df, table_name, file_name, sheet)
            total_loaded += rows_loaded
            
        except Exception as e:
            logger.error(f"Error processing {file_name} [{sheet}]: {e}")
            continue
            
    return total_loaded

def main():
    """Main function to load all planilha_embriologia Excel files to bronze tables"""
    logger.info("Starting Planilha Embriologia data loader (2021-2026)")
    logger.info(f"DuckDB path: {DUCKDB_PATH}")
    logger.info(f"Data input directory: {DATA_INPUT_DIR}")
    
    try:
        logger.info("Creating DuckDB connection...")
        con = get_duckdb_connection()
        logger.info("DuckDB connection created successfully")
        
        # Ensure excluded and obsolete tables are cleaned up from bronze schema
        con.execute("DROP TABLE IF EXISTS bronze.planilha_2022_bsb_sheet1")
        con.execute("DROP TABLE IF EXISTS bronze.planilha_2022_ibi_total")
        con.execute("DROP TABLE IF EXISTS bronze.planilha_2023_ibi_total_2023")
        con.execute("DROP TABLE IF EXISTS bronze.planilha_2023_ibi_total_2023_nova")
        con.execute("DROP TABLE IF EXISTS bronze.planilha_2023_ibira_total_2023")
        
        excel_files = get_all_excel_files()
        total_rows = 0
        files_processed = 0
        
        for file_path in excel_files:
            try:
                rows = process_excel_file(file_path, con)
                total_rows += rows
                files_processed += 1
            except Exception as e:
                logger.error(f"Failed to process {os.path.basename(file_path)}: {e}")
                continue
        
        # Final summary
        logger.info("=" * 50)
        logger.info("LOADING SUMMARY")
        logger.info("=" * 50)
        logger.info(f"Files processed: {files_processed}/{len(excel_files)}")
        logger.info(f"Total rows loaded to bronze: {total_rows:,}")
        logger.info("=" * 50)
        
        con.close()
        logger.info("Planilha Embriologia data loader completed")
        
    except Exception as e:
        logger.error(f"Error in main function: {e}")
        raise

if __name__ == "__main__":
    main()

