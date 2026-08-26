"""
Extract EmbryoScope Data for Cases to Evaluate (All Units - Year 2025)
======================================================================
Reads patients and PINs from 'Relação casos sem vídeo e foto.xlsx', matches
against huntington_data_lake.duckdb (silver_embryoscope) across all clinic units,
filtered strictly for embryos in YEAR 2025, extracts embryo_id prefix (up to 'P-'),
computes counts, and exports to Excel.
"""

import os
import re
import duckdb
import openpyxl
import pandas as pd
from datetime import datetime

# ──────────────────────────────────────────────────────────────
# File Paths
# ──────────────────────────────────────────────────────────────
SCRIPT_DIR = r"g:\My Drive\projetos_individuais\Huntington\db_recovery\casos_para_avaliar"
PROJECT_ROOT = r"g:\My Drive\projetos_individuais\Huntington"

INPUT_XLSX = os.path.join(SCRIPT_DIR, "data_input", "Relação casos sem vídeo e foto.xlsx")
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "data_output")
OUTPUT_XLSX = os.path.join(OUTPUT_DIR, "resultado_casos_avaliados.xlsx")
DB_PATH = os.path.join(PROJECT_ROOT, "database", "huntington_data_lake.duckdb")


def is_yellow_fill(cell) -> bool:
    """Check if an openpyxl cell has a yellow background fill."""
    try:
        if not cell.fill or not cell.fill.fill_type:
            return False
        if hasattr(cell.fill, "start_color") and cell.fill.start_color:
            sc = cell.fill.start_color
            if hasattr(sc, "rgb") and sc.rgb:
                rgb = str(sc.rgb).upper()
                if "FFFF00" in rgb or "FFFFE0" in rgb or rgb == "FFFFFF00":
                    return True
            if hasattr(sc, "index") and sc.index == 5:
                return True
    except Exception:
        pass
    return False


def load_input_patients(filepath: str) -> list[dict]:
    """Parse the input Excel file and extract patients and associated doctor groups."""
    print(f"Loading input file: {filepath}")
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb["Nomes"]
    
    current_doctor = "Unknown"
    patients = []
    
    for r in range(1, ws.max_row + 1):
        name_cell = ws.cell(row=r, column=1)
        name_val = str(name_cell.value).strip() if name_cell.value is not None else ""
        
        pin_cell = ws.cell(row=r, column=7)
        pin_val = str(pin_cell.value).strip() if pin_cell.value is not None else ""
        
        notes_cell = ws.cell(row=r, column=9)
        notes_val = str(notes_cell.value).strip() if notes_cell.value is not None else ""
        
        # Skip empty lines
        if not name_val and not pin_val:
            continue
            
        # Check if this is a header or section label
        if name_val.lower().startswith("pacientes fora") or name_val.lower() == "médicos":
            continue
            
        # Check if this is a yellow header row without a PIN (Doctor group)
        if is_yellow_fill(name_cell) and (not pin_val or pin_val.lower() == "none" or pin_val.lower() == "pin"):
            current_doctor = name_val
            print(f"  [Row {r:3d}] Doctor section: {current_doctor}")
            continue
            
        # If it's the header row ('Michelle' / 'PIN')
        if pin_val.lower() == "pin":
            if is_yellow_fill(name_cell):
                current_doctor = name_val
            continue
            
        # Patient row
        patients.append({
            "row_idx": r,
            "doctor_group": current_doctor,
            "input_patient_name": name_val,
            "input_pin": pin_val,
            "input_notes": notes_val
        })
        
    print(f"Loaded {len(patients)} patient rows from sheet 'Nomes'.")
    return patients


def query_embryo_data(patients: list[dict], db_path: str) -> list[dict]:
    """Match patients across all units, extract 2025 embryos, prefix, and counts."""
    print(f"Connecting to database: {db_path}")
    con = duckdb.connect(db_path, read_only=True, config={"access_mode": "READ_ONLY"})
    
    results = []
    
    for p in patients:
        pin = p["input_pin"]
        name = p["input_patient_name"]
        
        # 1. Match by prontuario across all locations
        matched_patients = con.execute("""
            SELECT PatientIDx, PatientID, FirstName, LastName, prontuario, _location
            FROM silver_embryoscope.patients
            WHERE CAST(prontuario AS VARCHAR) = ?
        """, [pin]).fetchall()
        match_type = "prontuario"
        
        # 2. Fallback: Match by PatientID across all locations
        if not matched_patients:
            matched_patients = con.execute("""
                SELECT PatientIDx, PatientID, FirstName, LastName, prontuario, _location
                FROM silver_embryoscope.patients
                WHERE CAST(PatientID AS VARCHAR) = ?
            """, [pin]).fetchall()
            if matched_patients:
                match_type = "patient_id"
                
        if not matched_patients:
            results.append({
                "doctor_group": p["doctor_group"],
                "input_patient_name": name,
                "input_pin": pin,
                "input_notes": p["input_notes"],
                "matched_patient_name": None,
                "unit_location": None,
                "patient_idx": None,
                "match_type": "None",
                "treatment_name": None,
                "embryo_id_prefix": None,
                "embryo_count": 0,
                "status": "Patient not found in any unit"
            })
            continue
            
        # Collect all 2025 embryos across all matching PatientIDx records for this patient
        all_patient_results = []
        
        for p_row in matched_patients:
            pidx, pat_id, first_n, last_n, pront, loc = p_row
            db_patient_name = f"{first_n or ''} {last_n or ''}".strip()
            
            # Fetch embryos for this patient and location
            embryos = con.execute("""
                SELECT EmbryoID, TreatmentName
                FROM silver_embryoscope.embryo_data
                WHERE _location = ? AND PatientIDx = ?
            """, [loc, pidx]).fetchall()
            
            # Filter strictly for 2025 embryos (EmbryoID pattern: D2025.MM.DD_...)
            embryos_2025 = [e for e in embryos if str(e[0]).startswith("D2025.")]
            
            if embryos_2025:
                # Group embryos by (TreatmentName, Prefix)
                prefix_groups = {}
                for embryo_id, treat_name in embryos_2025:
                    m = re.match(r"^(.*?P-)", str(embryo_id))
                    prefix = m.group(1) if m else str(embryo_id)
                    
                    key = (treat_name, prefix)
                    prefix_groups.setdefault(key, []).append(embryo_id)
                    
                for (treat_name, prefix), embryo_list in prefix_groups.items():
                    all_patient_results.append({
                        "doctor_group": p["doctor_group"],
                        "input_patient_name": name,
                        "input_pin": pin,
                        "input_notes": p["input_notes"],
                        "matched_patient_name": db_patient_name,
                        "unit_location": loc,
                        "patient_idx": pidx,
                        "match_type": match_type,
                        "treatment_name": treat_name,
                        "embryo_id_prefix": prefix,
                        "embryo_count": len(embryo_list),
                        "status": "Matched"
                    })
                    
        if all_patient_results:
            results.extend(all_patient_results)
        else:
            # None of the PatientIDx records had 2025 embryos
            p_first = matched_patients[0]
            db_name_first = f"{p_first[2] or ''} {p_first[3] or ''}".strip()
            results.append({
                "doctor_group": p["doctor_group"],
                "input_patient_name": name,
                "input_pin": pin,
                "input_notes": p["input_notes"],
                "matched_patient_name": db_name_first,
                "unit_location": p_first[5],
                "patient_idx": p_first[0],
                "match_type": match_type,
                "treatment_name": None,
                "embryo_id_prefix": None,
                "embryo_count": 0,
                "status": "No 2025 embryos found"
            })
                
    return results


def export_results(results: list[dict], output_path: str):
    """Save the results DataFrame to Excel and CSV files."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    df = pd.DataFrame(results)
    
    # Determine target Excel filename
    actual_xlsx = output_path
    try:
        with open(actual_xlsx, "a"):
            pass
    except (PermissionError, IOError):
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        actual_xlsx = output_path.replace(".xlsx", f"_2025_{ts}.xlsx")
        print(f"Warning: Primary Excel file is locked. Saving to fallback: {actual_xlsx}")

    print(f"Exporting results to Excel: {actual_xlsx}")
    with pd.ExcelWriter(actual_xlsx, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Resultados_2025", index=False)
        
        # Summary sheets
        matched_df = df[df["status"] == "Matched"]
        unmatched_df = df[df["status"] != "Matched"]
        
        matched_df.to_excel(writer, sheet_name="Casos_Identificados", index=False)
        unmatched_df.to_excel(writer, sheet_name="Casos_Nao_Identificados", index=False)
        
    # Save CSV backup
    csv_path = output_path.replace(".xlsx", ".csv")
    try:
        df.to_csv(csv_path, index=False, encoding="utf-8-sig")
        print(f"Exported CSV copy: {csv_path}")
    except Exception:
        csv_path = actual_xlsx.replace(".xlsx", ".csv")
        df.to_csv(csv_path, index=False, encoding="utf-8-sig")
        print(f"Exported CSV copy to fallback: {csv_path}")
    
    print("\n" + "=" * 60)
    print("EXTRACTION SUMMARY (ALL CLINICS - YEAR 2025)")
    print("=" * 60)
    print(f"Total output rows: {len(df)}")
    print(f"Unique input PINs evaluated: {df['input_pin'].nunique()}")
    print("\nBreakdown by Status:")
    print(df["status"].value_counts().to_string())
    print("\nBreakdown by Unit / Clinic Location:")
    print(df["unit_location"].value_counts(dropna=False).to_string())
    print("\nBreakdown by Match Type:")
    print(df["match_type"].value_counts().to_string())
    print("\nTotal 2025 Embryos identified across all prefixes:", df["embryo_count"].sum())
    print("=" * 60)


def main():
    patients = load_input_patients(INPUT_XLSX)
    results = query_embryo_data(patients, DB_PATH)
    export_results(results, OUTPUT_XLSX)


if __name__ == "__main__":
    main()
