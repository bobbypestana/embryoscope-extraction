#!/usr/bin/env python3
"""
REDLARA Bronze Ingestion Pipeline
Reads all Excel files from redlara/data_input/{unidade}/{year}/ folders and loads them into DuckDB bronze tables.
Supports all 12 workbooks (Ibirapuera, Santa Joana, Vila Mariana for 2021-2024).
Ingests all procedure sheets (FRESH, FET, FOT/FTO, OD/OR/RECEP, FP, IUI).
Includes streak-based termination to avoid parsing Excel empty formatted rows (e.g. redlara_vm_2022.xlsx).
Fully overwrites bronze.redlara_* tables on each run.
"""

import logging
import os
import re
import sys
import glob
from datetime import datetime
from pathlib import Path
import duckdb
import openpyxl
import pandas as pd

# Setup paths
SCRIPT_DIR = Path(__file__).resolve().parent
REDLARA_DIR = SCRIPT_DIR.parent
PROJECT_ROOT = REDLARA_DIR.parent
DATA_INPUT_DIR = REDLARA_DIR / "data_input"
DB_PATH = PROJECT_ROOT / "database" / "huntington_data_lake.duckdb"
LOGS_DIR = SCRIPT_DIR / "logs"
LOGS_DIR.mkdir(parents=True, exist_ok=True)

# Setup logging
timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
log_file = LOGS_DIR / f"01_redlara_to_bronze_{timestamp}.log"
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(str(log_file), encoding="utf-8"),
        logging.StreamHandler(sys.stdout)
    ]
)
logger = logging.getLogger("redlara_bronze")

def get_metadata_from_path(file_path: Path):
    """
    Extracts year and standardized unidade code and name.
    """
    filename = file_path.name
    path_str = str(file_path)
    
    # Extract Year
    year_match = re.search(r'(20\d{2})', path_str)
    year = int(year_match.group(1)) if year_match else None
    
    # Extract Unidade
    path_lower = path_str.lower()
    if 'ibirapuera' in path_lower or 'ibira' in path_lower:
        unidade_code = 'ibira'
        unidade_name = 'Ibirapuera'
    elif 'santa joana' in path_lower or '_sj' in path_lower or 'sj' in path_lower:
        unidade_code = 'sj'
        unidade_name = 'Santa Joana'
    elif 'vila mariana' in path_lower or '_vm' in path_lower or 'vm' in path_lower:
        unidade_code = 'vm'
        unidade_name = 'Vila Mariana'
    else:
        unidade_code = 'other'
        unidade_name = file_path.parent.parent.name
        
    return year, unidade_code, unidade_name

def classify_sheet(sheet_name: str) -> str:
    """
    Classifies Excel sheet name into one of the 6 standard procedure streams:
    fresh, fet, fot, recep, fp, iui. Returns None for non-data sheets (e.g. Folha1).
    """
    s = sheet_name.strip().upper()
    if s in ['FOLHA1', 'SHEET1', 'PLAN1']:
        return None
    
    # Priority classification
    if s.startswith('RECEP') or s.startswith('OD') or s == 'OR' or 'OD (OR)' in s:
        return 'recep'
    if s.startswith('FTO') or s.startswith('FOT'):
        return 'fot'
    if s.startswith('FRESH') or s.startswith('FIV'):
        return 'fresh'
    if s.startswith('FET') or s.startswith('TEC'):
        return 'fet'
    if s.startswith('FP'):
        return 'fp'
    if s.startswith('IUI') or s.startswith('IIU'):
        return 'iui'
        
    return None

def read_sheet_safely(file_path: Path, sheet_name: str):
    """
    Safely reads an Excel sheet using openpyxl in read_only mode.
    Stops after encountering 30 consecutive empty rows to prevent parsing
    up to 1,048,576 formatted empty cells in Excel workbooks.
    Detects dynamic header position and returns a clean DataFrame.
    """
    wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return None
        
    ws = wb[sheet_name]
    rows = []
    consecutive_empty = 0
    header_idx = None
    
    for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
        vals = [c for c in row if c is not None and str(c).strip() != '']
        if not vals:
            consecutive_empty += 1
            if header_idx is not None and consecutive_empty > 30:
                break
            continue
            
        consecutive_empty = 0
        
        # Detect header row if not found yet
        if header_idx is None:
            row_str = " ".join([str(v).lower() for v in vals])
            if any(k in row_str for k in ['outcome', 'procedure', 'chart', 'pin', 'date', 'nome', 'resultado', 'patient', 'diagnosis']):
                header_idx = len(rows)
                
        rows.append(list(row))
        
    wb.close()
    
    if not rows:
        return None
        
    if header_idx is None:
        header_idx = 0
        
    raw_headers = rows[header_idx]
    # Truncate trailing empty columns (e.g. Santa Joana 2022 FET has 16,384 columns with trailing blanks)
    non_empty_indices = [i for i, c in enumerate(raw_headers) if c is not None and str(c).strip() != '']
    max_col_idx = max(non_empty_indices) if non_empty_indices else len(raw_headers) - 1
    raw_headers = raw_headers[:max_col_idx + 1]

    headers = []
    for i, c in enumerate(raw_headers):
        if c is not None and str(c).strip() != '':
            headers.append(str(c).strip())
        else:
            headers.append(f"unnamed_col_{i}")
            
    data_rows = rows[header_idx + 1:]
    if not data_rows:
        return None
        
    cleaned_rows = []
    for r in data_rows:
        if len(r) < len(headers):
            r = list(r) + [None] * (len(headers) - len(r))
        else:
            r = list(r[:len(headers)])
            
        # Check if row has any non-empty value
        if any(c is not None and str(c).strip() != '' for c in r):
            cleaned_rows.append(r)
            
    if not cleaned_rows:
        return None
        
    df = pd.DataFrame(cleaned_rows, columns=headers)
    return df

def sanitize_column_name(col_name: str, used_names: set) -> str:
    """Sanitizes raw column names to valid, unique DuckDB column identifiers."""
    clean = str(col_name).strip()
    if not clean or clean.lower() == 'none' or clean.startswith('unnamed_col_'):
        clean = "col"
    
    # Remove newlines and non-alphanumeric chars
    clean = re.sub(r'[\r\n\t]+', '_', clean)
    clean = re.sub(r'[^a-zA-Z0-9_]', '_', clean)
    clean = re.sub(r'_+', '_', clean).strip('_').lower()
    if not clean:
        clean = "col"
        
    candidate = clean
    counter = 1
    while candidate in used_names:
        candidate = f"{clean}_{counter}"
        counter += 1
        
    used_names.add(candidate)
    return candidate

def ingest_all_bronze():
    """Main ingestion orchestrator."""
    logger.info("=" * 70)
    logger.info("Starting REDLARA Bronze Ingestion")
    logger.info(f"Target DuckDB: {DB_PATH}")
    logger.info(f"Data Input Directory: {DATA_INPUT_DIR}")
    logger.info("=" * 70)
    
    import time
    conn = None
    for attempt in range(1, 9):
        try:
            conn = duckdb.connect(str(DB_PATH))
            break
        except Exception as e:
            if attempt < 8:
                logger.warning(f"DuckDB connection attempt {attempt} failed ({e}). Retrying in 2s...")
                time.sleep(2)
            else:
                logger.error(f"DuckDB connection failed after 8 attempts: {e}")
                raise
                
    conn.execute("CREATE SCHEMA IF NOT EXISTS bronze")
    
    # 1. Cleanup existing bronze.redlara_* tables
    logger.info("Cleaning up existing bronze.redlara_% tables...")
    existing = conn.execute(
        "SELECT table_name FROM information_schema.tables WHERE table_schema='bronze' AND table_name LIKE 'redlara_%'"
    ).fetchall()
    for row in existing:
        conn.execute(f"DROP TABLE IF EXISTS bronze.{row[0]}")
    logger.info(f"Dropped {len(existing)} existing bronze.redlara_* tables.")
    
    # 2. Find all Excel files
    excel_files = sorted(list(DATA_INPUT_DIR.rglob("*.xlsx")))
    logger.info(f"Found {len(excel_files)} Excel files in data_input.")
    
    total_tables_created = 0
    total_rows_ingested = 0
    extraction_ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    for i, file_path in enumerate(excel_files, 1):
        year, unidade_code, unidade_name = get_metadata_from_path(file_path)
        rel_path = file_path.relative_to(DATA_INPUT_DIR)
        logger.info(f"[{i}/{len(excel_files)}] Processing file: {rel_path} (Unidade: {unidade_name}, Year: {year})")
        
        try:
            wb_peek = openpyxl.load_workbook(str(file_path), read_only=True)
            sheet_names = wb_peek.sheetnames
            wb_peek.close()
        except Exception as e:
            logger.error(f"  Failed to inspect workbook {file_path.name}: {e}")
            continue
            
        for sheet_name in sheet_names:
            procedure_stream = classify_sheet(sheet_name)
            if not procedure_stream:
                logger.info(f"  Skipping non-procedure sheet: '{sheet_name}'")
                continue
                
            table_name = f"redlara_{unidade_code}_{year}_{procedure_stream}"
            logger.info(f"  Reading sheet '{sheet_name}' -> target: bronze.{table_name}...")
            
            try:
                df = read_sheet_safely(file_path, sheet_name)
                if df is None or len(df) == 0:
                    logger.warning(f"  Sheet '{sheet_name}' yielded 0 rows. Skipping table creation.")
                    continue
                    
                # Clean column headers
                used_names = set()
                col_rename_map = {}
                for col in df.columns:
                    safe_col = sanitize_column_name(col, used_names)
                    col_rename_map[col] = safe_col
                    
                df = df.rename(columns=col_rename_map)
                
                # Cast all data columns to string and replace 'nan'/'None'
                df = df.astype(str)
                for c in df.columns:
                    df[c] = df[c].replace({'nan': None, 'None': None, '<NA>': None, '': None})
                    
                # Add lineage metadata without fragmentation
                df = df.copy()
                df = df.assign(
                    line_number=list(range(1, len(df) + 1)),
                    extraction_timestamp=extraction_ts,
                    file_name=file_path.name,
                    sheet_name=sheet_name,
                    unidade=unidade_name,
                    year=int(year) if year else None,
                    procedure_stream=procedure_stream
                )
                
                # Write to DuckDB
                conn.register('df_bronze_view', df)
                conn.execute(f"CREATE TABLE bronze.{table_name} AS SELECT * FROM df_bronze_view")
                conn.unregister('df_bronze_view')
                
                row_count = len(df)
                col_count = len(df.columns)
                total_tables_created += 1
                total_rows_ingested += row_count
                logger.info(f"  -> Created bronze.{table_name}: {row_count} rows, {col_count} cols")
                
            except Exception as e:
                logger.error(f"  Error processing sheet '{sheet_name}' in {file_path.name}: {e}", exc_info=True)
                
    conn.close()
    logger.info("=" * 70)
    logger.info("REDLARA Bronze Ingestion Completed Successfully")
    logger.info(f"Total Bronze Tables Created: {total_tables_created}")
    logger.info(f"Total Bronze Rows Ingested: {total_rows_ingested}")
    logger.info("=" * 70)

if __name__ == "__main__":
    ingest_all_bronze()
